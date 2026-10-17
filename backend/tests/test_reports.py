"""Module 6 - Reporting Agent: content builders, PDF/XLSX rendering, generator and endpoints."""

from collections.abc import Callable
from datetime import UTC, datetime
from decimal import Decimal
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.analytics.calculations import (
    MarkRecord,
    StudentInfo,
    SubjectConfig,
    at_risk_students,
    class_summary,
    student_results,
    subject_summary,
)
from app.analytics.loader import ExamRef
from app.analytics.service import ExamOverview, exam_overview
from app.config import settings
from app.dependencies import TenantContext
from app.exceptions import DomainValidationError, NotFoundError
from app.models import (
    AgentRun,
    AgentRunKind,
    AuditLog,
    InterventionPlan,
    InterventionStatus,
    MembershipRole,
    Report,
    ReportFormat,
    ReportKind,
    School,
    Teacher,
)
from app.reports.content import (
    FOOTER_TEXT,
    PlanRow,
    build_principal_content,
)
from app.reports.generator import generate_report, resolve_report_path, slugify
from app.reports.pdf import format_cell, safe_text
from tests.test_analytics_api import World, build_world

D = Decimal
BASE = "/api/v1/reports"
NOW = datetime(2026, 10, 16, 9, 30, tzinfo=UTC)


@pytest.fixture(autouse=True)
def reports_dir(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    monkeypatch.setattr(settings, "REPORTS_DIR", str(tmp_path))
    return tmp_path


@pytest.fixture()
def admin(make_tenant: Callable[..., TenantContext]) -> TenantContext:
    return make_tenant(role=MembershipRole.school_admin)


@pytest.fixture()
def world(db: Session, admin: TenantContext) -> World:
    return build_world(db, admin.school_id)


def _gen(db: Session, world: World, kind: ReportKind, fmt: ReportFormat, **kw: int) -> Report:
    extra: dict[str, int] = dict(kw)
    exam_id = extra.pop("exam_id", world.exam2)
    if kind == ReportKind.class_teacher:
        extra.setdefault("class_section_id", world.sec_a)
    if kind == ReportKind.subject_teacher:
        extra.setdefault("teacher_id", world.teacher)
    return generate_report(
        db, school_id=world.school_id, kind=kind, fmt=fmt, exam_id=exam_id, **extra
    )


def _sheet_values(path: Path, sheet: str) -> list[list[object]]:
    wb = load_workbook(path)
    return [list(row) for row in wb[sheet].iter_rows(values_only=True)]


def _find_row(rows: list[list[object]], first: object) -> list[object]:
    return next(r for r in rows if r and r[0] == first)


# ---------------------------------------------------------------------------
# Pure content builder
# ---------------------------------------------------------------------------


def _hand_overview() -> ExamOverview:
    configs = [
        SubjectConfig(1, "Maths", D(100), D(35)),
        SubjectConfig(2, "English", D(100), D(35)),
    ]
    students = [StudentInfo(i, f"S-{i:04d}", f"Kid {i}", 10, "8-A") for i in range(1, 5)]
    marks = [
        MarkRecord(1, 1, D(80)),
        MarkRecord(1, 2, D(90)),
        MarkRecord(2, 1, D(30)),  # fails maths
        MarkRecord(2, 2, D(70)),
        MarkRecord(3, 1, D(60)),
        MarkRecord(3, 2, D(60)),
        MarkRecord(4, 1, D(50)),
        MarkRecord(4, 2, D(40)),
    ]
    results = student_results(students, configs, marks)
    exam = ExamRef(1, "Unit Test 2", 2, None, 1, "2026-27")
    return ExamOverview(
        exam=exam,
        previous_exam=None,
        comparison_exam_id=None,
        class_section_ids=None,
        subjects=configs,
        school_summary=class_summary(results),
        class_summaries=[class_summary(results, 10, "8-A")],
        subject_summaries=subject_summary(results, configs),
        at_risk=at_risk_students(results, []),
        results=results,
        changes=[],
    )


def test_principal_narrative_from_hand_made_overview() -> None:
    overview = _hand_overview()
    plans = [
        PlanRow(1, 2, "S-0002", "Kid 2", "8-A", "Maths", "draft", "Practice sums", None, 7),
        PlanRow(2, 2, "S-0002", "Kid 2", "8-A", None, "approved", "Weekly check-in", None, 7),
    ]
    content = build_principal_content(
        school_name="Test School", overview=overview, previous=None, plans=plans, generated_at=NOW
    )
    narrative = content.narrative
    assert narrative[0] == (
        "4 students are in scope for Unit Test 2; 3 passed every subject "
        "(pass rate 75.0%, mean score 60.0%)."
    )
    # Maths 75% pass, English 100% -> average 87.5%, gap 12.5
    assert (
        "Maths has the lowest pass rate (75.0%), 12.5 points below the average subject "
        "pass rate (87.5%)." in narrative
    )
    assert "English has the highest pass rate (100.0%)." in narrative
    assert any(line.startswith("1 student meet the at-risk rule") for line in narrative)
    assert (
        "2 intervention plans are linked to this exam: 1 approved or in progress, "
        "1 awaiting teacher approval." in narrative
    )
    assert content.footer == FOOTER_TEXT
    titles = [s.title for s in content.sections]
    assert "Calculation method" in titles
    assert content.sheets == [
        "Summary", "Classes", "Subjects", "Students", "At-Risk", "Interventions", "Method",
    ]  # fmt: skip
    comparison = content.sections[0]
    assert comparison.paragraphs == ["There is no earlier exam to compare with."]
    at_risk = next(s for s in content.sections if s.sheet == "At-Risk")
    risk_list = at_risk.tables[-1]
    assert risk_list.rows[0][0] == "S-0002"
    assert risk_list.rows[0][-1] == "1 draft, 1 approved"


def test_format_helpers() -> None:
    from app.reports.content import Column

    assert format_cell(None, Column("x", "percent")) == "-"
    assert format_cell(D("62.5"), Column("x", "percent")) == "62.50%"
    assert format_cell(D("-3"), Column("x", "delta")) == "-3.00"
    assert format_cell(4, Column("x", "delta")) == "+4"
    assert format_cell(D("80.00"), Column("x", "decimal")) == "80.00"
    assert format_cell(D("85.00"), Column("x", "number")) == "85"
    assert safe_text("A & B -> ₹ →") == "A &amp; B -&gt; Rs. -&gt;"
    assert slugify("Quarterly Exam (Term 1)!") == "quarterly-exam-term-1"


# ---------------------------------------------------------------------------
# Generator
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("kind", list(ReportKind))
@pytest.mark.parametrize("fmt", list(ReportFormat))
def test_generate_each_kind_and_format(
    db: Session, world: World, reports_dir: Path, kind: ReportKind, fmt: ReportFormat
) -> None:
    report = _gen(db, world, kind, fmt)
    assert report.id is not None
    assert report.file_path == f"{world.school_id}/{report.id}_{kind.value}_unit-test-2.{fmt.value}"
    path = reports_dir / report.file_path
    assert path.is_file() and path.stat().st_size > 1000
    if fmt == ReportFormat.pdf:
        assert path.read_bytes().startswith(b"%PDF")
    else:
        wb = load_workbook(path)
        assert "Summary" in wb.sheetnames and "Method" in wb.sheetnames
    audit = db.scalar(
        select(AuditLog).where(AuditLog.entity_type == "report", AuditLog.entity_id == report.id)
    )
    assert audit is not None
    assert audit.actor == "ReportingAgent"
    assert audit.detail is not None and audit.detail["exam_id"] == world.exam2


def test_principal_xlsx_matches_analytics(db: Session, world: World, reports_dir: Path) -> None:
    report = _gen(db, world, ReportKind.principal, ReportFormat.xlsx)
    path = reports_dir / report.file_path
    wb = load_workbook(path)
    assert wb.sheetnames == [
        "Summary", "Classes", "Subjects", "Students", "At-Risk", "Interventions", "Method",
    ]  # fmt: skip
    ov = exam_overview(db, world.school_id, world.exam2)

    classes = _sheet_values(path, "Classes")
    for summary in ov.class_summaries:
        row = _find_row(classes, summary.class_label)
        assert row[1] == summary.student_count
        assert row[3] == summary.pass_count
        assert row[5] == pytest.approx(float(summary.pass_percentage or 0))
        assert row[6] == pytest.approx(float(summary.mean_percentage or 0))

    subjects = _sheet_values(path, "Subjects")
    for s in ov.subject_summaries:
        row = _find_row(subjects, s.subject_name)
        assert row[6] == pytest.approx(float(s.overall.pass_percentage or 0))

    students = _sheet_values(path, "Students")
    for r in ov.results:
        row = next(x for x in students if len(x) > 2 and x[2] == r.pseudo_id)
        assert row[-3] == pytest.approx(float(r.percentage))
        assert row[-2] == ("Pass" if r.passed else ("Absent" if not r.appeared else "Fail"))

    at_risk = _sheet_values(path, "At-Risk")
    listed = {row[0] for row in at_risk if isinstance(row[0], str) and row[0].startswith("S-")}
    assert listed == {a.pseudo_id for a in ov.at_risk}

    ws = wb["Students"]
    assert ws.freeze_panes is not None
    assert ws.auto_filter.ref is not None
    assert len(ws.conditional_formatting) >= 1
    summary_text = [c for row in _sheet_values(path, "Summary") for c in row if c]
    assert FOOTER_TEXT in summary_text
    assert any(isinstance(c, str) and "pass rate" in c for c in summary_text)


def test_class_teacher_xlsx_only_has_that_class(
    db: Session, world: World, reports_dir: Path
) -> None:
    report = _gen(db, world, ReportKind.class_teacher, ReportFormat.xlsx)
    students = _sheet_values(reports_dir / report.file_path, "Students")
    ids = {row[1] for row in students if isinstance(row[1], str) and row[1].startswith("S-")}
    assert ids == {"S-0001", "S-0002", "S-0003", "S-0004"}


def test_subject_teacher_xlsx_has_coaching_view(
    db: Session, world: World, reports_dir: Path
) -> None:
    report = _gen(db, world, ReportKind.subject_teacher, ReportFormat.xlsx)
    path = reports_dir / report.file_path
    assert load_workbook(path).sheetnames == ["Summary", "Classes", "Students", "Method"]
    classes = _sheet_values(path, "Classes")
    labelled = [r for r in classes if r[0] in {"8-A", "8-B"}]
    assert {(r[0], r[1]) for r in labelled} == {
        ("8-A", "Maths"),
        ("8-B", "Maths"),
        ("8-A", "English"),
    }
    flat = " ".join(str(c) for row in classes for c in row if c)
    assert "rank" not in flat.lower().replace("not a rating or ranking", "")


def test_plans_linked_to_exam_appear(db: Session, world: World, reports_dir: Path) -> None:
    run = AgentRun(
        school_id=world.school_id, kind=AgentRunKind.full_analysis, input={"exam_id": world.exam2}
    )
    other = AgentRun(
        school_id=world.school_id, kind=AgentRunKind.full_analysis, input={"exam_id": world.exam1}
    )
    db.add_all([run, other])
    db.flush()
    db.add_all(
        [
            InterventionPlan(
                school_id=world.school_id, run_id=run.id, student_id=world.students[4],
                subject_id=world.maths, status=InterventionStatus.draft, plan_text="Daily practice",
            ),
            InterventionPlan(
                school_id=world.school_id, run_id=run.id, student_id=world.students[6],
                status=InterventionStatus.approved, plan_text="Weekly check-in",
            ),
            InterventionPlan(
                school_id=world.school_id, run_id=other.id, student_id=world.students[5],
                status=InterventionStatus.draft, plan_text="Old exam plan",
            ),
        ]
    )  # fmt: skip
    db.flush()
    report = _gen(db, world, ReportKind.principal, ReportFormat.xlsx, run_id=run.id)
    assert report.run_id == run.id
    rows = _sheet_values(reports_dir / report.file_path, "Interventions")
    assert _find_row(rows, "Draft (awaiting approval)")[1] == 1
    assert _find_row(rows, "Approved")[1] == 1
    assert _find_row(rows, "Total")[1] == 2
    texts = {c for row in rows for c in row if isinstance(c, str)}
    assert "Daily practice" in texts and "Old exam plan" not in texts


def test_generator_invalid_combos_and_foreign_ids(
    db: Session, world: World, make_tenant: Callable[..., TenantContext]
) -> None:
    with pytest.raises(DomainValidationError):
        generate_report(
            db, school_id=world.school_id, kind=ReportKind.class_teacher,
            fmt=ReportFormat.pdf, exam_id=world.exam2,
        )  # fmt: skip
    with pytest.raises(DomainValidationError):
        generate_report(
            db, school_id=world.school_id, kind=ReportKind.subject_teacher,
            fmt=ReportFormat.pdf, exam_id=world.exam2,
        )  # fmt: skip
    other = make_tenant()
    foreign = build_world(db, other.school_id)
    with pytest.raises(NotFoundError):
        _gen(db, world, ReportKind.principal, ReportFormat.pdf, exam_id=foreign.exam2)
    with pytest.raises(NotFoundError):
        _gen(db, world, ReportKind.class_teacher, ReportFormat.pdf, class_section_id=foreign.sec_a)
    with pytest.raises(NotFoundError):
        _gen(db, world, ReportKind.subject_teacher, ReportFormat.pdf, teacher_id=foreign.teacher)
    run = AgentRun(school_id=other.school_id, kind=AgentRunKind.report)
    db.add(run)
    db.flush()
    with pytest.raises(NotFoundError):
        _gen(db, world, ReportKind.principal, ReportFormat.pdf, run_id=run.id)


def test_empty_exam_still_renders(db: Session, world: World, reports_dir: Path) -> None:
    for fmt in ReportFormat:
        report = _gen(db, world, ReportKind.principal, fmt, exam_id=world.exam3_empty)
        assert (reports_dir / report.file_path).stat().st_size > 500


def test_path_traversal_rejected(db: Session, world: World) -> None:
    report = Report(
        school_id=world.school_id, kind=ReportKind.principal, format=ReportFormat.pdf,
        file_path="../../etc/passwd",
    )  # fmt: skip
    with pytest.raises(NotFoundError):
        resolve_report_path(report)
    report.file_path = f"{world.school_id + 1}/1_principal_x.pdf"
    with pytest.raises(NotFoundError):
        resolve_report_path(report)


# ---------------------------------------------------------------------------
# Endpoints
# ---------------------------------------------------------------------------


def test_endpoint_flow(
    db: Session,
    world: World,
    admin: TenantContext,
    client_as: Callable[[TenantContext], TestClient],
    reports_dir: Path,
) -> None:
    client = client_as(admin)
    res = client.post(BASE, json={"kind": "principal", "format": "pdf", "exam_id": world.exam2})
    assert res.status_code == 201, res.text
    created = res.json()
    assert created["exam_name"] == "Unit Test 2"
    assert created["scope_label"] == "Whole school"
    assert created["created_by"] == admin.user.id
    assert created["created_by_name"] == admin.user.full_name
    assert created["download_url"] == f"/api/v1/reports/{created['id']}/download"

    res = client.post(
        BASE,
        json={
            "kind": "class_teacher", "format": "xlsx",
            "exam_id": world.exam1, "class_section_id": world.sec_b,
        },
    )  # fmt: skip
    assert res.status_code == 201, res.text
    second = res.json()
    assert second["scope_label"] == "Class 8-B"

    listed = client.get(BASE).json()
    assert [r["id"] for r in listed] == [second["id"], created["id"]]
    assert [r["id"] for r in client.get(BASE, params={"exam_id": world.exam2}).json()] == [
        created["id"]
    ]
    assert [r["id"] for r in client.get(BASE, params={"kind": "class_teacher"}).json()] == [
        second["id"]
    ]
    assert client.get(BASE, params={"run_id": 999999}).json() == []

    dl = client.get(f"{BASE}/{created['id']}/download")
    assert dl.status_code == 200
    assert dl.headers["content-type"] == "application/pdf"
    assert created["file_name"] in dl.headers["content-disposition"]
    assert dl.content.startswith(b"%PDF")
    dl = client.get(f"{BASE}/{second['id']}/download")
    assert dl.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    path = reports_dir / str(world.school_id) / created["file_name"]
    assert path.is_file()
    assert client.delete(f"{BASE}/{created['id']}").status_code == 204
    assert not path.exists()
    assert client.get(f"{BASE}/{created['id']}/download").status_code == 404
    assert db.get(Report, created["id"]) is None


def test_invalid_combos_return_422(
    world: World, admin: TenantContext, client_as: Callable[[TenantContext], TestClient]
) -> None:
    client = client_as(admin)
    for body in (
        {"kind": "class_teacher", "format": "pdf", "exam_id": world.exam2},
        {"kind": "subject_teacher", "format": "xlsx", "exam_id": world.exam2},
        {"kind": "nope", "format": "pdf", "exam_id": world.exam2},
        {"kind": "principal", "format": "docx", "exam_id": world.exam2},
    ):
        res = client.post(BASE, json=body)
        assert res.status_code == 422, body
        assert res.json()["error"]["code"] == "VALIDATION_ERROR"


def test_delete_requires_school_admin(
    db: Session,
    world: World,
    make_tenant: Callable[..., TenantContext],
    client_as: Callable[[TenantContext], TestClient],
) -> None:
    report = _gen(db, world, ReportKind.principal, ReportFormat.pdf)
    school = db.get_one(School, world.school_id)
    principal = make_tenant(role=MembershipRole.principal, school=school)
    client = client_as(principal)
    assert client.get(f"{BASE}/{report.id}/download").status_code == 200
    assert client.delete(f"{BASE}/{report.id}").status_code == 403


def test_cross_tenant_returns_404(
    db: Session,
    world: World,
    make_tenant: Callable[..., TenantContext],
    client_as: Callable[[TenantContext], TestClient],
) -> None:
    report = _gen(db, world, ReportKind.principal, ReportFormat.pdf)
    outsider = make_tenant(role=MembershipRole.school_admin)
    client = client_as(outsider)
    assert client.get(f"{BASE}/{report.id}/download").status_code == 404
    assert client.delete(f"{BASE}/{report.id}").status_code == 404
    assert client.get(BASE).json() == []
    res = client.post(BASE, json={"kind": "principal", "format": "pdf", "exam_id": world.exam2})
    assert res.status_code == 404


def test_teacher_permissions(
    db: Session,
    world: World,
    admin: TenantContext,
    make_tenant: Callable[..., TenantContext],
    client_as: Callable[[TenantContext], TestClient],
) -> None:
    school = db.get_one(School, world.school_id)
    teacher_ctx = make_tenant(role=MembershipRole.teacher, school=school)
    rao = db.get_one(Teacher, world.teacher)
    rao.user_id = teacher_ctx.user.id
    db.flush()
    # A report about another scope generated by the admin - invisible to the teacher.
    hidden = _gen(db, world, ReportKind.principal, ReportFormat.pdf)
    # A report about the teacher's own class generated by the admin - visible.
    shared = _gen(
        db, world, ReportKind.class_teacher, ReportFormat.pdf, class_section_id=world.sec_a
    )
    db.commit()

    client = client_as(teacher_ctx)
    exam = world.exam2
    assert client.post(BASE, json={"kind": "principal", "exam_id": exam}).status_code == 403
    own = client.post(
        BASE, json={"kind": "subject_teacher", "exam_id": exam, "teacher_id": world.teacher}
    )
    assert own.status_code == 201, own.text
    other_teacher = client.post(
        BASE,
        json={"kind": "subject_teacher", "exam_id": exam, "teacher_id": world.teacher_unassigned},
    )
    assert other_teacher.status_code == 403
    taught = client.post(
        BASE, json={"kind": "class_teacher", "exam_id": exam, "class_section_id": world.sec_b}
    )
    assert taught.status_code == 201, taught.text
    not_taught = client.post(
        BASE,
        json={"kind": "class_teacher", "exam_id": exam, "class_section_id": world.sec_c_empty},
    )
    assert not_taught.status_code == 403

    visible = {r["id"] for r in client.get(BASE).json()}
    assert visible == {own.json()["id"], taught.json()["id"], shared.id}
    assert client.get(f"{BASE}/{hidden.id}/download").status_code == 404
    assert client.get(f"{BASE}/{shared.id}/download").status_code == 200
    assert client.delete(f"{BASE}/{own.json()['id']}").status_code == 403
