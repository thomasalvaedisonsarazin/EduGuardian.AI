"""Excel renderer (openpyxl).

One sheet per ``ReportContent.sheets`` entry (Summary, Classes, Subjects, Students, At-Risk,
Interventions, Method). Tables get bold filled headers, number formats, column widths, and on
data sheets the first table gets freeze panes + an autofilter. "Status" columns carry
conditional fills (red for Fail/Absent/High, green for Pass) - the text stays, so meaning does
not rely on colour alone.
"""

from __future__ import annotations

import logging
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.reports.content import (
    NEGATIVE_STATUSES,
    POSITIVE_STATUSES,
    SHEET_SUMMARY,
    WARNING_STATUSES,
    CellValue,
    ReportContent,
    Section,
    Table,
)

logger = logging.getLogger(__name__)

__all__ = ["render_xlsx"]

TITLE_FONT = Font(bold=True, size=14, color="312E81")
SECTION_FONT = Font(bold=True, size=12, color="312E81")
TABLE_TITLE_FONT = Font(bold=True, size=11)
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="4F46E5")
MUTED_FONT = Font(italic=True, color="64748B", size=9)
FAIL_FILL = PatternFill("solid", fgColor="FEE2E2", bgColor="FEE2E2")
PASS_FILL = PatternFill("solid", fgColor="DCFCE7", bgColor="DCFCE7")
WARN_FILL = PatternFill("solid", fgColor="FEF3C7", bgColor="FEF3C7")
FLAG_FONT = Font(bold=True, color="991B1B")
WRAP = Alignment(wrap_text=True, vertical="top")

NUMBER_FORMATS = {
    "int": "0",
    "number": "0.##",
    "decimal": "0.00",
    "percent": '0.00"%"',
    "delta": "+0.00;-0.00;0.00",
}
MIN_WIDTH = 8
MAX_WIDTH = 60


def _excel_value(value: CellValue) -> str | int | float | None:
    return float(value) if isinstance(value, Decimal) else value


def _quoted(values: frozenset[str]) -> list[str]:
    return [f'"{v}"' for v in sorted(values)]


class _SheetWriter:
    def __init__(self, ws: Worksheet, is_summary: bool) -> None:
        self.ws = ws
        self.row = 1
        self.is_summary = is_summary
        self.filtered = False
        self.widths: dict[int, int] = {}

    def _track(self, col: int, text: str) -> None:
        longest = max((len(part) for part in text.split("\n")), default=0)
        self.widths[col] = max(self.widths.get(col, MIN_WIDTH), min(longest + 2, MAX_WIDTH))

    def text(self, value: str, font: Font | None = None) -> None:
        cell = self.ws.cell(row=self.row, column=1, value=value)
        if font is not None:
            cell.font = font
        self.row += 1

    def blank(self) -> None:
        self.row += 1

    def section(self, section: Section) -> None:
        self.text(section.title, SECTION_FONT)
        for paragraph in section.paragraphs:
            self.text(paragraph)
        for bullet in section.bullets:
            self.text(f"- {bullet}")
        for table in section.tables:
            self.blank()
            self.table(table)
        if section.chart is not None and section.chart.values:
            self.blank()
            self.text(f"{section.chart.title} (chart data)", TABLE_TITLE_FONT)
            for label, value in zip(section.chart.labels, section.chart.values, strict=True):
                self.ws.cell(row=self.row, column=1, value=label)
                c = self.ws.cell(row=self.row, column=2, value=value)
                c.number_format = NUMBER_FORMATS["percent"]
                self._track(1, label)
                self.row += 1
        self.blank()

    def table(self, table: Table) -> None:
        ws = self.ws
        self.text(table.title, TABLE_TITLE_FONT)
        if not table.rows:
            self.text(table.empty_message, MUTED_FONT)
            return
        header_row = self.row
        for c_idx, column in enumerate(table.columns, start=1):
            cell = ws.cell(row=header_row, column=c_idx, value=column.header)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            self._track(c_idx, column.header)
        for r_off, row in enumerate(table.rows):
            excel_row = header_row + 1 + r_off
            for c_idx, (value, column) in enumerate(zip(row, table.columns, strict=True), 1):
                cell = ws.cell(row=excel_row, column=c_idx, value=_excel_value(value))
                if column.kind in NUMBER_FORMATS and not isinstance(value, str):
                    cell.number_format = NUMBER_FORMATS[column.kind]
                if (r_off, c_idx - 1) in table.flagged:
                    cell.font = FLAG_FONT
                    cell.fill = FAIL_FILL
                text = "" if value is None else str(value)
                if len(text) > MAX_WIDTH:
                    cell.alignment = WRAP
                self._track(c_idx, text)
        last_row = header_row + len(table.rows)
        for c_idx, column in enumerate(table.columns, start=1):
            if column.kind != "status":
                continue
            letter = get_column_letter(c_idx)
            rng = f"{letter}{header_row + 1}:{letter}{last_row}"
            first = f"{letter}{header_row + 1}"
            for values, fill in (
                (NEGATIVE_STATUSES, FAIL_FILL),
                (POSITIVE_STATUSES, PASS_FILL),
                (WARNING_STATUSES, WARN_FILL),
            ):
                checks = ",".join(f"{first}={v}" for v in _quoted(values))
                ws.conditional_formatting.add(
                    rng, FormulaRule(formula=[f"OR({checks})"], fill=fill, stopIfTrue=True)
                )
        if not self.is_summary and not self.filtered:
            ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
            ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(table.columns))}{last_row}"
            self.filtered = True
        self.row = last_row + 1
        if table.note:
            self.text(table.note, MUTED_FONT)

    def finish(self) -> None:
        for col, width in self.widths.items():
            self.ws.column_dimensions[get_column_letter(col)].width = width


def render_xlsx(content: ReportContent, path: Path) -> None:
    """Write ``content`` as an .xlsx workbook to ``path``."""
    wb = Workbook()
    default = wb.active
    writers: dict[str, _SheetWriter] = {}
    for idx, name in enumerate(content.sheets):
        ws = default if idx == 0 and default is not None else wb.create_sheet(name)
        ws.title = name
        writers[name] = _SheetWriter(ws, is_summary=name == SHEET_SUMMARY)

    summary = writers.get(SHEET_SUMMARY) or next(iter(writers.values()))
    summary.text(content.title, TITLE_FONT)
    summary.text(f"{content.school_name} · {content.subtitle}")
    summary.text(f"Generated {content.generated_at.strftime('%Y-%m-%d %H:%M')}", MUTED_FONT)
    summary.blank()
    if content.figures:
        summary.text("Key figures", SECTION_FONT)
        for fig in content.figures:
            summary.ws.cell(row=summary.row, column=1, value=fig.label).font = Font(bold=True)
            summary.ws.cell(row=summary.row, column=2, value=fig.value)
            if fig.hint:
                summary.ws.cell(row=summary.row, column=3, value=fig.hint).font = MUTED_FONT
            summary._track(1, fig.label)
            summary._track(2, fig.value)
            summary.row += 1
        summary.blank()
    if content.narrative:
        summary.text("Key findings", SECTION_FONT)
        for line in content.narrative:
            summary.text(f"- {line}")
        summary.blank()

    for section in content.sections:
        writer = writers.get(section.sheet, summary)
        writer.section(section)

    for name, writer in writers.items():
        if writer.row == 1:
            writer.text(f"No {name.lower()} data for this report.", MUTED_FONT)
        writer.blank()
        writer.text(content.footer, MUTED_FONT)
        writer.finish()
    # Narrative lines are long; keep column A readable rather than sized to the longest one.
    summary.ws.column_dimensions["A"].width = max(
        summary.ws.column_dimensions["A"].width or MIN_WIDTH, 28
    )
    wb.properties.title = content.title
    wb.properties.creator = "EduGuardian AI"
    wb.save(path)
    logger.info("XLSX report written: %s (%d bytes)", path, path.stat().st_size)
