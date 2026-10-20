"""Business logic for Module 2 (Academic Setup).

Every function takes the caller's ``school_id`` and filters by it; ids received from clients are
resolved through ``_get_owned`` so another school's object always surfaces as a 404.
Deletes use bulk ``DELETE`` statements so the database's ``ON DELETE CASCADE`` rules apply.
"""

import csv
import io
import itertools
import logging
import re
from collections.abc import Iterable, Sequence
from dataclasses import dataclass
from typing import Any

from sqlalchemy import delete, func, or_, select, update
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.config import settings
from app.exceptions import ConflictError, DomainValidationError, NotFoundError
from app.models import (
    AcademicYear,
    ClassSection,
    Exam,
    ExamSubjectConfig,
    Mark,
    Membership,
    Student,
    Subject,
    Teacher,
    TeachingAssignment,
)
from app.schemas.setup import (
    AcademicYearCreate,
    AcademicYearUpdate,
    ClassSectionCreate,
    ClassSectionOut,
    ClassSectionUpdate,
    ExamCreate,
    ExamSubjectConfigIn,
    ExamSubjectConfigOut,
    ExamUpdate,
    ImportRowError,
    StudentCreate,
    StudentImportResult,
    StudentUpdate,
    SubjectCreate,
    SubjectUpdate,
    TeacherCreate,
    TeacherUpdate,
    TeachingAssignmentCreate,
    TeachingAssignmentUpdate,
)
from app.services.audit import log_action

logger = logging.getLogger(__name__)

PSEUDO_ID_PREFIX = "S-"
_PSEUDO_RE = re.compile(r"^S-(\d+)$")
IMPORT_REQUIRED_COLUMNS = ("admission_no", "full_name", "grade", "section")


@dataclass(frozen=True)
class Actor:
    """Who performs a write (for audit logging)."""

    school_id: int
    user_id: int | None


# =============================================================================================
# helpers
# =============================================================================================
def _get_owned[ModelT](
    db: Session, model: type[ModelT], obj_id: int, school_id: int, name: str
) -> ModelT:
    obj = db.get(model, obj_id)
    if obj is None or getattr(obj, "school_id", None) != school_id:
        raise NotFoundError(name)
    return obj


def _audit(
    db: Session, actor: Actor, action: str, entity_type: str, entity_id: int | None, **detail: Any
) -> None:
    log_action(
        db,
        school_id=actor.school_id,
        actor="user",
        action=action,
        entity_type=entity_type,
        entity_id=entity_id,
        user_id=actor.user_id,
        detail={k: _jsonable(v) for k, v in detail.items()},
    )


def _jsonable(value: Any) -> Any:
    if isinstance(value, dict):
        return {k: _jsonable(v) for k, v in value.items()}
    if isinstance(value, list | tuple):
        return [_jsonable(v) for v in value]
    if value is None or isinstance(value, bool | int | float | str):
        return value
    return str(value)


def _commit(db: Session, conflict_message: str) -> None:
    """Commit, translating constraint races into a 409."""
    try:
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        logger.info("Integrity error on commit: %s", exc.orig)
        raise ConflictError(conflict_message) from exc


def _apply(obj: object, changes: dict[str, Any], nullable: Iterable[str] = ()) -> None:
    allowed_null = set(nullable)
    for key, value in changes.items():
        if value is None and key not in allowed_null:
            continue
        setattr(obj, key, value)


def _count_marks(db: Session, *conditions: Any) -> int:
    return int(db.scalar(select(func.count(Mark.id)).where(*conditions)) or 0)


# =============================================================================================
# academic years
# =============================================================================================
def list_academic_years(db: Session, school_id: int) -> Sequence[AcademicYear]:
    stmt = (
        select(AcademicYear)
        .where(AcademicYear.school_id == school_id)
        .order_by(AcademicYear.label.desc())
    )
    return db.scalars(stmt).all()


def get_academic_year(db: Session, school_id: int, year_id: int) -> AcademicYear:
    return _get_owned(db, AcademicYear, year_id, school_id, "Academic year")


def get_current_academic_year(db: Session, school_id: int) -> AcademicYear | None:
    return db.scalars(
        select(AcademicYear).where(
            AcademicYear.school_id == school_id, AcademicYear.is_current.is_(True)
        )
    ).first()


def _ensure_unique_year_label(
    db: Session, school_id: int, label: str, exclude_id: int | None = None
) -> None:
    stmt = select(AcademicYear.id).where(
        AcademicYear.school_id == school_id, AcademicYear.label == label
    )
    if exclude_id is not None:
        stmt = stmt.where(AcademicYear.id != exclude_id)
    if db.scalar(stmt) is not None:
        raise ConflictError(f"Academic year '{label}' already exists")


def _unset_other_current(db: Session, school_id: int, keep_id: int) -> None:
    db.execute(
        update(AcademicYear)
        .where(AcademicYear.school_id == school_id, AcademicYear.id != keep_id)
        .values(is_current=False)
    )


def create_academic_year(db: Session, actor: Actor, data: AcademicYearCreate) -> AcademicYear:
    _ensure_unique_year_label(db, actor.school_id, data.label)
    year = AcademicYear(school_id=actor.school_id, label=data.label, is_current=data.is_current)
    db.add(year)
    db.flush()
    if year.is_current:
        _unset_other_current(db, actor.school_id, year.id)
    _audit(db, actor, "academic_year.create", "academic_year", year.id, **data.model_dump())
    _commit(db, f"Academic year '{data.label}' already exists")
    db.refresh(year)
    return year


def update_academic_year(
    db: Session, actor: Actor, year_id: int, data: AcademicYearUpdate
) -> AcademicYear:
    year = get_academic_year(db, actor.school_id, year_id)
    changes = data.model_dump(exclude_unset=True)
    if changes.get("label"):
        changes["label"] = changes["label"].strip()
        _ensure_unique_year_label(db, actor.school_id, changes["label"], exclude_id=year.id)
    _apply(year, changes)
    db.flush()
    if year.is_current:
        _unset_other_current(db, actor.school_id, year.id)
    _audit(db, actor, "academic_year.update", "academic_year", year.id, **changes)
    _commit(db, "Academic year label already exists")
    db.refresh(year)
    return year


def delete_academic_year(db: Session, actor: Actor, year_id: int) -> None:
    year = get_academic_year(db, actor.school_id, year_id)
    exam_ids = select(Exam.id).where(Exam.academic_year_id == year.id)
    if _count_marks(db, Mark.exam_id.in_(exam_ids)):
        raise ConflictError(
            f"Academic year '{year.label}' has exams with recorded marks and cannot be deleted"
        )
    _audit(db, actor, "academic_year.delete", "academic_year", year.id, label=year.label)
    db.execute(delete(AcademicYear).where(AcademicYear.id == year.id))
    db.commit()
    db.expire_all()


# =============================================================================================
# class sections
# =============================================================================================
def _class_out(cs: ClassSection, student_count: int) -> ClassSectionOut:
    return ClassSectionOut(
        id=cs.id,
        school_id=cs.school_id,
        academic_year_id=cs.academic_year_id,
        grade=cs.grade,
        section=cs.section,
        student_count=student_count,
    )


def _student_counts(db: Session, class_ids: Sequence[int]) -> dict[int, int]:
    if not class_ids:
        return {}
    rows = db.execute(
        select(Student.class_section_id, func.count(Student.id))
        .where(Student.class_section_id.in_(class_ids))
        .group_by(Student.class_section_id)
    ).all()
    return {int(cid): int(n) for cid, n in rows}


def list_class_sections(
    db: Session, school_id: int, academic_year_id: int | None = None
) -> list[ClassSectionOut]:
    stmt = select(ClassSection).where(ClassSection.school_id == school_id)
    if academic_year_id is not None:
        stmt = stmt.where(ClassSection.academic_year_id == academic_year_id)
    stmt = stmt.order_by(ClassSection.grade, ClassSection.section)
    classes = db.scalars(stmt).all()
    counts = _student_counts(db, [c.id for c in classes])
    return [_class_out(c, counts.get(c.id, 0)) for c in classes]


def get_class_section(db: Session, school_id: int, class_id: int) -> ClassSection:
    return _get_owned(db, ClassSection, class_id, school_id, "Class")


def get_class_section_out(db: Session, school_id: int, class_id: int) -> ClassSectionOut:
    cs = get_class_section(db, school_id, class_id)
    return _class_out(cs, _student_counts(db, [cs.id]).get(cs.id, 0))


def _ensure_unique_class(
    db: Session, school_id: int, year_id: int, grade: int, section: str, exclude_id: int | None
) -> None:
    stmt = select(ClassSection.id).where(
        ClassSection.school_id == school_id,
        ClassSection.academic_year_id == year_id,
        ClassSection.grade == grade,
        ClassSection.section == section,
    )
    if exclude_id is not None:
        stmt = stmt.where(ClassSection.id != exclude_id)
    if db.scalar(stmt) is not None:
        raise ConflictError(f"Class {grade}-{section} already exists in this academic year")


def create_class_section(db: Session, actor: Actor, data: ClassSectionCreate) -> ClassSectionOut:
    get_academic_year(db, actor.school_id, data.academic_year_id)
    _ensure_unique_class(db, actor.school_id, data.academic_year_id, data.grade, data.section, None)
    cs = ClassSection(school_id=actor.school_id, **data.model_dump())
    db.add(cs)
    db.flush()
    _audit(db, actor, "class_section.create", "class_section", cs.id, **data.model_dump())
    _commit(db, "Class already exists")
    return _class_out(cs, 0)


def update_class_section(
    db: Session, actor: Actor, class_id: int, data: ClassSectionUpdate
) -> ClassSectionOut:
    cs = get_class_section(db, actor.school_id, class_id)
    changes = data.model_dump(exclude_unset=True)
    if changes.get("academic_year_id") is not None:
        get_academic_year(db, actor.school_id, changes["academic_year_id"])
    _apply(cs, changes)
    _ensure_unique_class(db, actor.school_id, cs.academic_year_id, cs.grade, cs.section, cs.id)
    db.flush()
    _audit(db, actor, "class_section.update", "class_section", cs.id, **changes)
    _commit(db, "Class already exists")
    return get_class_section_out(db, actor.school_id, cs.id)


def delete_class_section(db: Session, actor: Actor, class_id: int) -> None:
    cs = get_class_section(db, actor.school_id, class_id)
    student_ids = select(Student.id).where(Student.class_section_id == cs.id)
    if _count_marks(db, Mark.student_id.in_(student_ids)):
        raise ConflictError(
            f"Class {cs.grade}-{cs.section} has students with recorded marks and cannot be deleted"
        )
    _audit(
        db,
        actor,
        "class_section.delete",
        "class_section",
        cs.id,
        grade=cs.grade,
        section=cs.section,
    )
    db.execute(delete(ClassSection).where(ClassSection.id == cs.id))
    db.commit()
    db.expire_all()


# =============================================================================================
# subjects
# =============================================================================================
def list_subjects(db: Session, school_id: int) -> Sequence[Subject]:
    return db.scalars(
        select(Subject).where(Subject.school_id == school_id).order_by(Subject.name)
    ).all()


def get_subject(db: Session, school_id: int, subject_id: int) -> Subject:
    return _get_owned(db, Subject, subject_id, school_id, "Subject")


def _ensure_unique_subject_code(
    db: Session, school_id: int, code: str, exclude_id: int | None
) -> None:
    stmt = select(Subject.id).where(Subject.school_id == school_id, Subject.code == code)
    if exclude_id is not None:
        stmt = stmt.where(Subject.id != exclude_id)
    if db.scalar(stmt) is not None:
        raise ConflictError(f"Subject code '{code}' already exists")


def create_subject(db: Session, actor: Actor, data: SubjectCreate) -> Subject:
    _ensure_unique_subject_code(db, actor.school_id, data.code, None)
    subject = Subject(school_id=actor.school_id, name=data.name, code=data.code)
    db.add(subject)
    db.flush()
    _audit(db, actor, "subject.create", "subject", subject.id, **data.model_dump())
    _commit(db, f"Subject code '{data.code}' already exists")
    return subject


def update_subject(db: Session, actor: Actor, subject_id: int, data: SubjectUpdate) -> Subject:
    subject = get_subject(db, actor.school_id, subject_id)
    changes: dict[str, Any] = {
        k: v.strip() if isinstance(v, str) else v
        for k, v in data.model_dump(exclude_unset=True).items()
    }
    if changes.get("code"):
        _ensure_unique_subject_code(db, actor.school_id, changes["code"], subject.id)
    _apply(subject, changes)
    db.flush()
    _audit(db, actor, "subject.update", "subject", subject.id, **changes)
    _commit(db, "Subject code already exists")
    return subject


def delete_subject(db: Session, actor: Actor, subject_id: int) -> None:
    subject = get_subject(db, actor.school_id, subject_id)
    if _count_marks(db, Mark.subject_id == subject.id):
        raise ConflictError(f"Subject '{subject.name}' has recorded marks and cannot be deleted")
    _audit(db, actor, "subject.delete", "subject", subject.id, name=subject.name)
    db.execute(delete(Subject).where(Subject.id == subject.id))
    db.commit()
    db.expire_all()


# =============================================================================================
# teachers
# =============================================================================================
def list_teachers(db: Session, school_id: int) -> Sequence[Teacher]:
    return db.scalars(
        select(Teacher).where(Teacher.school_id == school_id).order_by(Teacher.full_name)
    ).all()


def get_teacher(db: Session, school_id: int, teacher_id: int) -> Teacher:
    return _get_owned(db, Teacher, teacher_id, school_id, "Teacher")


def _ensure_member(db: Session, school_id: int, user_id: int) -> None:
    stmt = select(Membership.id).where(
        Membership.school_id == school_id, Membership.user_id == user_id
    )
    if db.scalar(stmt) is None:
        raise NotFoundError("User")


def create_teacher(db: Session, actor: Actor, data: TeacherCreate) -> Teacher:
    if data.user_id is not None:
        _ensure_member(db, actor.school_id, data.user_id)
    teacher = Teacher(
        school_id=actor.school_id, full_name=data.full_name.strip(), user_id=data.user_id
    )
    db.add(teacher)
    db.flush()
    _audit(db, actor, "teacher.create", "teacher", teacher.id, **data.model_dump())
    _commit(db, "Teacher could not be created")
    return teacher


def update_teacher(db: Session, actor: Actor, teacher_id: int, data: TeacherUpdate) -> Teacher:
    teacher = get_teacher(db, actor.school_id, teacher_id)
    changes = data.model_dump(exclude_unset=True)
    if changes.get("user_id") is not None:
        _ensure_member(db, actor.school_id, changes["user_id"])
    if changes.get("full_name"):
        changes["full_name"] = changes["full_name"].strip()
    _apply(teacher, changes, nullable=("user_id",))
    db.flush()
    _audit(db, actor, "teacher.update", "teacher", teacher.id, **changes)
    _commit(db, "Teacher could not be updated")
    return teacher


def delete_teacher(db: Session, actor: Actor, teacher_id: int) -> None:
    teacher = get_teacher(db, actor.school_id, teacher_id)
    _audit(db, actor, "teacher.delete", "teacher", teacher.id, full_name=teacher.full_name)
    db.execute(delete(Teacher).where(Teacher.id == teacher.id))
    db.commit()
    db.expire_all()


# =============================================================================================
# teaching assignments
# =============================================================================================
def list_teaching_assignments(
    db: Session,
    school_id: int,
    teacher_id: int | None = None,
    class_section_id: int | None = None,
) -> Sequence[TeachingAssignment]:
    stmt = select(TeachingAssignment).where(TeachingAssignment.school_id == school_id)
    if teacher_id is not None:
        stmt = stmt.where(TeachingAssignment.teacher_id == teacher_id)
    if class_section_id is not None:
        stmt = stmt.where(TeachingAssignment.class_section_id == class_section_id)
    return db.scalars(stmt.order_by(TeachingAssignment.id)).all()


def get_teaching_assignment(db: Session, school_id: int, ta_id: int) -> TeachingAssignment:
    return _get_owned(db, TeachingAssignment, ta_id, school_id, "Teaching assignment")


def _validate_assignment_refs(
    db: Session, school_id: int, teacher_id: int, class_section_id: int, subject_id: int
) -> None:
    get_teacher(db, school_id, teacher_id)
    get_class_section(db, school_id, class_section_id)
    get_subject(db, school_id, subject_id)


def _ensure_unique_assignment(db: Session, ta: TeachingAssignment, exclude_id: int | None) -> None:
    stmt = select(TeachingAssignment.id).where(
        TeachingAssignment.teacher_id == ta.teacher_id,
        TeachingAssignment.class_section_id == ta.class_section_id,
        TeachingAssignment.subject_id == ta.subject_id,
    )
    if exclude_id is not None:
        stmt = stmt.where(TeachingAssignment.id != exclude_id)
    if db.scalar(stmt) is not None:
        raise ConflictError("This teacher is already assigned to that subject in that class")


def _enforce_single_class_teacher(db: Session, ta: TeachingAssignment) -> None:
    """A section has one class teacher: clear the flag on other teachers' assignments."""
    if not ta.is_class_teacher:
        return
    db.execute(
        update(TeachingAssignment)
        .where(
            TeachingAssignment.school_id == ta.school_id,
            TeachingAssignment.class_section_id == ta.class_section_id,
            TeachingAssignment.teacher_id != ta.teacher_id,
        )
        .values(is_class_teacher=False)
    )


def create_teaching_assignment(
    db: Session, actor: Actor, data: TeachingAssignmentCreate
) -> TeachingAssignment:
    _validate_assignment_refs(
        db, actor.school_id, data.teacher_id, data.class_section_id, data.subject_id
    )
    ta = TeachingAssignment(school_id=actor.school_id, **data.model_dump())
    _ensure_unique_assignment(db, ta, None)
    db.add(ta)
    db.flush()
    _enforce_single_class_teacher(db, ta)
    _audit(
        db, actor, "teaching_assignment.create", "teaching_assignment", ta.id, **data.model_dump()
    )
    _commit(db, "Duplicate teaching assignment")
    db.refresh(ta)
    return ta


def update_teaching_assignment(
    db: Session, actor: Actor, ta_id: int, data: TeachingAssignmentUpdate
) -> TeachingAssignment:
    ta = get_teaching_assignment(db, actor.school_id, ta_id)
    changes = data.model_dump(exclude_unset=True)
    _apply(ta, changes)
    _validate_assignment_refs(
        db, actor.school_id, ta.teacher_id, ta.class_section_id, ta.subject_id
    )
    _ensure_unique_assignment(db, ta, ta.id)
    db.flush()
    _enforce_single_class_teacher(db, ta)
    _audit(db, actor, "teaching_assignment.update", "teaching_assignment", ta.id, **changes)
    _commit(db, "Duplicate teaching assignment")
    db.refresh(ta)
    return ta


def delete_teaching_assignment(db: Session, actor: Actor, ta_id: int) -> None:
    ta = get_teaching_assignment(db, actor.school_id, ta_id)
    _audit(db, actor, "teaching_assignment.delete", "teaching_assignment", ta.id)
    db.execute(delete(TeachingAssignment).where(TeachingAssignment.id == ta.id))
    db.commit()
    db.expire_all()


# =============================================================================================
# students
# =============================================================================================
def _max_pseudo_number(db: Session, school_id: int) -> int:
    ids = db.scalars(
        select(Student.pseudo_id).where(
            Student.school_id == school_id, Student.pseudo_id.like(f"{PSEUDO_ID_PREFIX}%")
        )
    ).all()
    best = 0
    for pid in ids:
        m = _PSEUDO_RE.match(pid)
        if m:
            best = max(best, int(m.group(1)))
    return best


def format_pseudo_id(number: int) -> str:
    return f"{PSEUDO_ID_PREFIX}{number:04d}"


def next_pseudo_id(db: Session, school_id: int) -> str:
    """Next per-school pseudonymous id: ``S-0001``, ``S-0002``, … (max existing + 1)."""
    return format_pseudo_id(_max_pseudo_number(db, school_id) + 1)


def list_students(
    db: Session,
    school_id: int,
    class_section_id: int | None = None,
    q: str | None = None,
    limit: int = 50,
    offset: int = 0,
) -> tuple[Sequence[Student], int]:
    conditions: list[Any] = [Student.school_id == school_id]
    if class_section_id is not None:
        conditions.append(Student.class_section_id == class_section_id)
    if q and q.strip():
        pattern = f"%{q.strip()}%"
        conditions.append(
            or_(
                Student.full_name.ilike(pattern),
                Student.admission_no.ilike(pattern),
                Student.pseudo_id.ilike(pattern),
            )
        )
    total = int(db.scalar(select(func.count(Student.id)).where(*conditions)) or 0)
    items = db.scalars(
        select(Student)
        .where(*conditions)
        .order_by(Student.full_name, Student.id)
        .limit(limit)
        .offset(offset)
    ).all()
    return items, total


def get_student(db: Session, school_id: int, student_id: int) -> Student:
    return _get_owned(db, Student, student_id, school_id, "Student")


def _ensure_unique_admission(
    db: Session, school_id: int, admission_no: str, exclude_id: int | None
) -> None:
    stmt = select(Student.id).where(
        Student.school_id == school_id, Student.admission_no == admission_no
    )
    if exclude_id is not None:
        stmt = stmt.where(Student.id != exclude_id)
    if db.scalar(stmt) is not None:
        raise ConflictError(f"Admission number '{admission_no}' already exists")


def create_student(db: Session, actor: Actor, data: StudentCreate) -> Student:
    get_class_section(db, actor.school_id, data.class_section_id)
    _ensure_unique_admission(db, actor.school_id, data.admission_no, None)
    student = Student(
        school_id=actor.school_id,
        class_section_id=data.class_section_id,
        admission_no=data.admission_no,
        full_name=data.full_name,
        pseudo_id=next_pseudo_id(db, actor.school_id),
    )
    db.add(student)
    db.flush()
    _audit(
        db,
        actor,
        "student.create",
        "student",
        student.id,
        class_section_id=data.class_section_id,
        pseudo_id=student.pseudo_id,
    )
    _commit(db, "Student could not be created (duplicate admission number?)")
    return student


def update_student(db: Session, actor: Actor, student_id: int, data: StudentUpdate) -> Student:
    student = get_student(db, actor.school_id, student_id)
    changes: dict[str, Any] = {
        k: v.strip() if isinstance(v, str) else v
        for k, v in data.model_dump(exclude_unset=True).items()
    }
    if changes.get("class_section_id") is not None:
        get_class_section(db, actor.school_id, changes["class_section_id"])
    if changes.get("admission_no"):
        _ensure_unique_admission(db, actor.school_id, changes["admission_no"], student.id)
    _apply(student, changes)
    db.flush()
    # Names are PII: audit which fields changed, not the values.
    _audit(db, actor, "student.update", "student", student.id, fields=sorted(changes))
    _commit(db, "Student could not be updated (duplicate admission number?)")
    return student


def delete_student(db: Session, actor: Actor, student_id: int) -> None:
    student = get_student(db, actor.school_id, student_id)
    if _count_marks(db, Mark.student_id == student.id):
        raise ConflictError(f"Student {student.pseudo_id} has recorded marks and cannot be deleted")
    _audit(db, actor, "student.delete", "student", student.id, pseudo_id=student.pseudo_id)
    db.execute(delete(Student).where(Student.id == student.id))
    db.commit()
    db.expire_all()


# --- student import ---------------------------------------------------------------------------
def _norm_header(value: object) -> str:
    return re.sub(r"[\s\-]+", "_", str(value or "").strip().lower())


def _cell_str(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def parse_tabular(content: bytes, filename: str) -> list[dict[str, str]]:
    """Parse CSV or XLSX bytes into rows keyed by normalised header (lowercase, snake_case)."""
    name = (filename or "").lower()
    max_rows = settings.MAX_IMPORT_ROWS
    raw_rows: list[list[object]]
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        try:
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:  # openpyxl raises many types for bad files
            raise DomainValidationError("Could not read the Excel file") from exc
        ws = wb.worksheets[0]
        raw_rows = [list(r) for r in ws.iter_rows(values_only=True, max_row=max_rows + 2)]
        wb.close()
    elif name.endswith(".csv") or not name:
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = content.decode("latin-1")
        raw_rows = [list(r) for r in itertools.islice(csv.reader(io.StringIO(text)), max_rows + 2)]
    else:
        raise DomainValidationError("Unsupported file type; upload a .csv or .xlsx file")

    if not raw_rows:
        raise DomainValidationError("The file is empty")
    if len(raw_rows) - 1 > max_rows:
        raise DomainValidationError(
            f"The file has more than {max_rows} rows - split it into smaller files",
            code="TOO_MANY_ROWS",
        )
    headers = [_norm_header(h) for h in raw_rows[0]]
    rows: list[dict[str, str]] = []
    for raw in raw_rows[1:]:
        cells = [_cell_str(c) for c in raw]
        rows.append({h: (cells[i] if i < len(cells) else "") for i, h in enumerate(headers) if h})
    return rows


def import_students(
    db: Session, actor: Actor, content: bytes, filename: str
) -> StudentImportResult:
    """Upsert students by admission_no. Rows with errors are skipped and reported."""
    rows = parse_tabular(content, filename)
    present = set(rows[0].keys()) if rows else set()
    if rows:
        missing = [c for c in IMPORT_REQUIRED_COLUMNS if c not in present]
        if missing:
            raise DomainValidationError(f"Missing required column(s): {', '.join(missing)}")

    year = get_current_academic_year(db, actor.school_id)
    if year is None:
        raise DomainValidationError("Set a current academic year before importing students")

    classes = {
        (c.grade, c.section.upper()): c.id
        for c in db.scalars(
            select(ClassSection).where(
                ClassSection.school_id == actor.school_id,
                ClassSection.academic_year_id == year.id,
            )
        )
    }
    existing = {
        s.admission_no: s
        for s in db.scalars(select(Student).where(Student.school_id == actor.school_id))
    }
    next_number = _max_pseudo_number(db, actor.school_id) + 1

    created = updated = 0
    errors: list[ImportRowError] = []
    seen: set[str] = set()
    for idx, row in enumerate(rows, start=2):  # row 1 is the header
        if not any(v for v in row.values()):
            continue
        admission_no = row.get("admission_no", "")
        full_name = row.get("full_name", "")
        grade_raw = row.get("grade", "")
        section = row.get("section", "").upper()
        problems: list[str] = []
        if not admission_no:
            problems.append("admission_no is required")
        elif len(admission_no) > 50:
            problems.append("admission_no is too long")
        if not full_name:
            problems.append("full_name is required")
        grade: int | None = None
        try:
            grade = int(float(grade_raw))
        except ValueError:
            problems.append(f"grade '{grade_raw}' is not a number")
        if not section:
            problems.append("section is required")
        class_id = classes.get((grade, section)) if grade is not None else None
        if grade is not None and section and class_id is None:
            problems.append(f"class {grade}-{section} does not exist in {year.label}")
        if admission_no and admission_no in seen:
            problems.append(f"duplicate admission_no '{admission_no}' in file")
        if problems:
            errors.append(ImportRowError(row=idx, message="; ".join(problems)))
            continue
        assert class_id is not None
        seen.add(admission_no)

        student = existing.get(admission_no)
        if student is None:
            student = Student(
                school_id=actor.school_id,
                class_section_id=class_id,
                admission_no=admission_no,
                full_name=full_name[:200],
                pseudo_id=format_pseudo_id(next_number),
            )
            next_number += 1
            db.add(student)
            existing[admission_no] = student
            created += 1
        else:
            student.full_name = full_name[:200]
            student.class_section_id = class_id
            updated += 1

    _audit(
        db,
        actor,
        "student.import",
        "student",
        None,
        filename=filename,
        created=created,
        updated=updated,
        errors=len(errors),
    )
    _commit(db, "Import failed due to a conflicting record; please retry")
    return StudentImportResult(created=created, updated=updated, errors=errors)


# =============================================================================================
# exams
# =============================================================================================
def list_exams(db: Session, school_id: int, academic_year_id: int | None = None) -> Sequence[Exam]:
    stmt = select(Exam).where(Exam.school_id == school_id)
    if academic_year_id is not None:
        stmt = stmt.where(Exam.academic_year_id == academic_year_id)
    return db.scalars(stmt.order_by(Exam.sequence, Exam.id)).all()


def get_exam(db: Session, school_id: int, exam_id: int) -> Exam:
    return _get_owned(db, Exam, exam_id, school_id, "Exam")


def create_exam(db: Session, actor: Actor, data: ExamCreate) -> Exam:
    get_academic_year(db, actor.school_id, data.academic_year_id)
    exam = Exam(school_id=actor.school_id, **data.model_dump())
    exam.name = exam.name.strip()
    db.add(exam)
    db.flush()
    _audit(db, actor, "exam.create", "exam", exam.id, **data.model_dump())
    _commit(db, "Exam could not be created")
    return exam


def update_exam(db: Session, actor: Actor, exam_id: int, data: ExamUpdate) -> Exam:
    exam = get_exam(db, actor.school_id, exam_id)
    changes = data.model_dump(exclude_unset=True)
    if changes.get("academic_year_id") is not None:
        get_academic_year(db, actor.school_id, changes["academic_year_id"])
    _apply(exam, changes, nullable=("held_on",))
    db.flush()
    _audit(db, actor, "exam.update", "exam", exam.id, **changes)
    _commit(db, "Exam could not be updated")
    return exam


def delete_exam(db: Session, actor: Actor, exam_id: int) -> None:
    exam = get_exam(db, actor.school_id, exam_id)
    if _count_marks(db, Mark.exam_id == exam.id):
        raise ConflictError(f"Exam '{exam.name}' has recorded marks and cannot be deleted")
    _audit(db, actor, "exam.delete", "exam", exam.id, name=exam.name)
    db.execute(delete(Exam).where(Exam.id == exam.id))
    db.commit()
    db.expire_all()


def list_exam_subjects(db: Session, school_id: int, exam_id: int) -> list[ExamSubjectConfigOut]:
    exam = get_exam(db, school_id, exam_id)
    rows = db.execute(
        select(ExamSubjectConfig, Subject)
        .join(Subject, Subject.id == ExamSubjectConfig.subject_id)
        .where(ExamSubjectConfig.exam_id == exam.id, ExamSubjectConfig.school_id == school_id)
        .order_by(Subject.name)
    ).all()
    return [
        ExamSubjectConfigOut(
            id=cfg.id,
            school_id=cfg.school_id,
            exam_id=cfg.exam_id,
            subject_id=cfg.subject_id,
            subject_name=subj.name,
            subject_code=subj.code,
            max_marks=cfg.max_marks,
            pass_marks=cfg.pass_marks,
        )
        for cfg, subj in rows
    ]


def replace_exam_subjects(
    db: Session, actor: Actor, exam_id: int, configs: Sequence[ExamSubjectConfigIn]
) -> list[ExamSubjectConfigOut]:
    """Replace the exam's subject configs with ``configs`` (upsert + remove the rest)."""
    exam = get_exam(db, actor.school_id, exam_id)
    subject_ids = [c.subject_id for c in configs]
    if len(set(subject_ids)) != len(subject_ids):
        raise DomainValidationError("Each subject can appear only once")
    for sid in subject_ids:
        get_subject(db, actor.school_id, sid)

    current = {
        cfg.subject_id: cfg
        for cfg in db.scalars(select(ExamSubjectConfig).where(ExamSubjectConfig.exam_id == exam.id))
    }
    removed = [sid for sid in current if sid not in set(subject_ids)]
    if removed and _count_marks(db, Mark.exam_id == exam.id, Mark.subject_id.in_(removed)):
        raise ConflictError(
            "Cannot remove a subject from this exam because marks are already recorded for it"
        )
    if removed:
        db.execute(
            delete(ExamSubjectConfig).where(
                ExamSubjectConfig.exam_id == exam.id, ExamSubjectConfig.subject_id.in_(removed)
            )
        )
    for c in configs:
        cfg = current.get(c.subject_id)
        if cfg is None:
            db.add(
                ExamSubjectConfig(
                    school_id=actor.school_id,
                    exam_id=exam.id,
                    subject_id=c.subject_id,
                    max_marks=c.max_marks,
                    pass_marks=c.pass_marks,
                )
            )
        else:
            cfg.max_marks = c.max_marks
            cfg.pass_marks = c.pass_marks
    _audit(
        db,
        actor,
        "exam.subjects.replace",
        "exam",
        exam.id,
        configs=[c.model_dump() for c in configs],
    )
    _commit(db, "Subject configuration conflict")
    db.expire_all()
    return list_exam_subjects(db, actor.school_id, exam.id)
